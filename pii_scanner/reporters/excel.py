from __future__ import annotations

import os
import re

import openpyxl

from pii_scanner.core.models import ScanResult, Status
from pii_scanner.reporters.budget import RowBudget
from pii_scanner.reporters.summary import Summary, summarize

# Excel 시트 행 한도(1,048,576) 미만으로 findings 행을 캡 — 초과분은 생략하고
# summary 시트에 생략 건수를 기록한다(전체 데이터는 state JSONL 이 원본).
MAX_FINDINGS_ROWS = 1_048_000

# xlsx(=XML)에 넣을 수 없는 제어문자와, 디코딩 불가 바이트에서 온 lone surrogate 를 지운다.
# 파일·폴더 이름에 이런 문자가 섞여 있으면 openpyxl 이 저장 단계에서 죽어 리포트가 한 장도
# 안 나온다 — 스캔은 다 해놓고 산출물을 통째로 잃는 실패라 여기서 막는다.
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _safe_text(v: str) -> str:
    return _ILLEGAL.sub("\ufffd", v).encode("utf-8", "replace").decode("utf-8")


def write_excel_stream(files, out_path: str, summary: Summary,
                       max_findings: int = MAX_FINDINGS_ROWS) -> None:
    """FileResult 이터러블을 1회 순회하며 report.xlsx 를 스트리밍으로 쓴다(T-014).

    대형 스캔(hit 4,300만/state 7.7GB)은 전체 결과를 RAM 이나 시트에
    담을 수 없다 — write_only 워크북으로 행 단위 append 만 하고, findings 가
    max_findings 를 넘으면 이후 행은 생략한다. 이터러블은 1회성(제너레이터)일 수
    있으므로 집계는 summarize_iter 로 미리 계산해 summary 로 받는다.
    """
    s = summary
    budget = RowBudget(s, max_findings)     # 고유식별정보·노출 → 고유식별정보·마스킹 → 나머지
    wb = openpyxl.Workbook(write_only=True)

    ws = wb.create_sheet("findings")
    ws.append([
        "파일경로", "파일명", "PII종류", "위험등급", "상태", "마스킹스니펫", "위치", "검증",
    ])
    we = wb.create_sheet("errors")
    # '구분' 이 없으면 "폴더를 못 열었다"와 "파일 추출이 실패했다"가 같은 줄로 보인다 —
    # 감사에서 이 둘은 의미가 전혀 다르므로 열로 갈라 둔다.
    we.append(["파일경로", "사유", "구분"])
    wenc = wb.create_sheet("encrypted")
    wenc.append(["파일경로"])

    for fr in files:
        if fr.unreadable:                  # 순회 실패 — 스캔조차 못 한 경로
            we.append([_safe_text(fr.path), _safe_text(fr.error or ""), "접근 실패"])
            continue
        if fr.encrypted:
            wenc.append([fr.path])
            continue                       # 암호화 파일은 hits 없음
        if fr.error:
            we.append([_safe_text(fr.path), _safe_text(fr.error), "파일 처리 실패"])
        for h in fr.hits:
            if not budget.allow(h):
                continue
            ws.append([
                fr.path,
                os.path.basename(fr.path),
                h.pii_type.value,
                h.risk.value,
                h.status.value,
                h.snippet,          # 마스킹 스니펫 — 원문 아님
                h.location or "",
                h.confidence.value,
            ])

    wsum = wb.create_sheet("summary")
    wsum.append(["항목", "값"])
    # 출처를 맨 위에 — 감사 산출물로 제출될 때 "언제·어느 버전으로 스캔했나"가 먼저 보여야 한다.
    if s.scanned_at:
        wsum.append(["스캔 일시", s.scanned_at])
    if s.tool_version:
        wsum.append(["도구 버전", s.tool_version])
    if s.active_types:
        # '이번 실행' 이라고 못박는 이유: 재개(Dropbox)에서는 리포트가 이전 실행 결과까지
        # 누적해 그리므로, 이 목록이 리포트 전체를 진술한다고 하면 거짓이 될 수 있다.
        wsum.append(["검사한 종류(이번 실행)", ", ".join(s.active_types)])
    if s.inactive_types:
        wsum.append(["⚠ 검사하지 않은 종류(이번 실행)", ", ".join(s.inactive_types)])
    wsum.append(["스캔 파일 수", s.total_files])
    wsum.append(["추출 실패 수", s.error_files])
    wsum.append(["암호화 건너뜀 수", s.encrypted_files])
    wsum.append(["노출(EXPOSED)", s.exposed])
    wsum.append(["마스킹(MASKED)", s.masked])
    wsum.append(["마스킹률(%)", round(s.masking_rate, 1)])
    wsum.append(["법인등록번호 오탐 제거", s.corp_filtered])
    if s.unreadable_paths:
        # 0일 때 행을 넣지 않는 이유: 순회 실패를 보고하는 커넥터는 로컬뿐이라, Dropbox
        # 리포트에 '0'을 찍으면 검증하지도 않은 커버리지를 보증하는 거짓 진술이 된다.
        wsum.append(["접근 실패(읽지 못한 폴더) 수 — errors 시트 '접근 실패' 참조",
                     s.unreadable_paths])
    if budget.skipped:
        wsum.append(["findings 생략 행수(시트 한도 초과 — 고유식별정보 우선 수록, "
                     "전체는 state JSONL 참조)", budget.skipped])
    wsum.append([])
    wsum.append(["PII종류", "노출", "마스킹"])
    for pii_type, b in s.by_type.items():
        wsum.append([pii_type.value, b["exposed"], b["masked"]])

    wb.save(out_path)


def write_excel(result: ScanResult, out_path: str, summary: Summary | None = None) -> None:
    """소형 스캔용 래퍼. summary 를 주면 그것을 쓴다 — 리포트 3종이 같은 집계·출처를 쓰도록."""
    write_excel_stream(result.files, out_path, summary or summarize(result))
