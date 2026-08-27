from __future__ import annotations

import io
import os

import docx
import openpyxl
import xlrd

from pii_scanner.core.extractors.base import Extractor, EncryptedFileError
from pii_scanner.core.locate import CellLocationIndex, Locator

# 비밀번호 걸린 OOXML 은 OLE 복합문서로 저장된다 — 이 매직으로 시작하면 암호화로 판정.
_OLE_MAGIC = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"


def _raise_if_encrypted(path: str) -> None:
    with open(path, "rb") as f:
        head = f.read(8)
    if head == _OLE_MAGIC:
        raise EncryptedFileError(f"암호화된 OOXML 추정(OLE 컨테이너): {os.path.basename(path)}")


class DocxExtractor(Extractor):
    def extract(self, path: str) -> str:
        _raise_if_encrypted(path)
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)


class XlsxExtractor(Extractor):
    def extract(self, path: str) -> str:
        return self.extract_located(path)[0]

    def extract_located(self, path: str) -> tuple[str, Locator]:
        _raise_if_encrypted(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        # 셀 조각을 리스트에, 라벨을 문자열로 쌓으면 대형 시트에서 메모리가 터진다(T-018:
        # 2,300만 셀 = 라벨만 4.5GB). 텍스트는 버퍼에 이어 쓰고 위치는 정수 배열로만 담는다.
        buf = io.StringIO()
        idx = CellLocationIndex()
        pos = 0
        for ws in wb.worksheets:
            idx.sheet(ws.title)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:     # EmptyCell(value=None)은 좌표 접근 전에 걸러짐
                        s = str(cell.value)
                        if pos:
                            buf.write("\n")        # 첫 셀 앞에는 안 붙임 → "\n".join 과 동일
                        idx.add(pos, cell.row, cell.column)
                        buf.write(s)
                        pos += len(s) + 1
        wb.close()
        return buf.getvalue(), idx.build()


class XlsExtractor(Extractor):
    """구 엑셀 .xls(BIFF) — xlrd 로 읽기 전용 텍스트 추출.

    .xls 는 정상도 OLE(D0CF11E0)라 _raise_if_encrypted(OOXML 용 매직 검사)를 쓰지 않는다.
    암호화는 xlrd 가 던지는 XLRDError('...encrypted...') 신호로만 판정한다(보수적).
    """

    def extract(self, path: str) -> str:
        return self.extract_located(path)[0]

    def extract_located(self, path: str) -> tuple[str, Locator]:
        try:
            book = xlrd.open_workbook(path)
        except xlrd.XLRDError as exc:
            if "encrypt" in str(exc).lower():
                raise EncryptedFileError(f"암호화된 xls: {path}") from exc
            raise                                  # 그 외 XLRDError 는 그대로 전파
        buf = io.StringIO()
        idx = CellLocationIndex()                  # xlsx 와 같은 이유(T-018)
        pos = 0
        for sheet in book.sheets():
            idx.sheet(sheet.name)
            for r in range(sheet.nrows):
                for c in range(sheet.ncols):
                    v = sheet.cell_value(r, c)
                    if v != "":
                        s = str(v)
                        if pos:
                            buf.write("\n")
                        idx.add(pos, r + 1, c + 1)   # xlrd 는 0-기준 → 1-기준으로
                        buf.write(s)
                        pos += len(s) + 1
        return buf.getvalue(), idx.build()
