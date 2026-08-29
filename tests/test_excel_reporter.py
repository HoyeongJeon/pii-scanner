import openpyxl
from pii_scanner.reporters.excel import write_excel
from pii_scanner.core.models import (
    ScanResult, FileResult, PiiHit, PiiType, Status, Confidence, RiskLevel,
)


def _result():
    sr = ScanResult()
    fr = FileResult(path="/x/a.txt")
    fr.hits.append(PiiHit(PiiType.RRN, Status.EXPOSED, "900101-1******", 0, 13,
                          Confidence.CONFIRMED, RiskLevel.CRITICAL,
                          location="Sheet1!C5"))
    sr.files.append(fr)
    sr.files.append(FileResult(path="/x/bad.txt", error="추출 실패: X"))
    return sr


def test_excel_has_four_sheets_and_no_raw(tmp_path):
    out = tmp_path / "r.xlsx"
    write_excel(_result(), str(out))
    wb = openpyxl.load_workbook(str(out))
    assert set(wb.sheetnames) == {"findings", "errors", "summary", "encrypted"}
    # findings 에 마스킹 스니펫만, 원문 없음
    findings_text = " ".join(
        str(c.value) for row in wb["findings"].iter_rows() for c in row if c.value
    )
    assert "900101-1******" in findings_text
    assert "1234568" not in findings_text   # 원문 절대 미기록


def test_excel_findings_include_location_column(tmp_path):
    out = tmp_path / "r.xlsx"
    write_excel(_result(), str(out))
    wb = openpyxl.load_workbook(str(out))
    rows = list(wb["findings"].iter_rows(values_only=True))
    header = rows[0]
    assert "위치" in header
    assert rows[1][header.index("위치")] == "Sheet1!C5"


def test_excel_stream_caps_findings_and_notes_skipped(tmp_path):
    # T-014: findings 가 상한(Excel 행 한도)을 넘으면 생략 + summary 에 생략 건수 기록
    from pii_scanner.reporters.excel import write_excel_stream
    from pii_scanner.reporters.summary import summarize
    sr = ScanResult()
    fr = FileResult(path="/x/many.txt")
    for i in range(5):
        fr.hits.append(PiiHit(PiiType.PHONE, Status.EXPOSED, f"010-****-{i:04d}",
                              0, 13, Confidence.CONFIRMED, RiskLevel.HIGH))
    sr.files.append(fr)
    out = tmp_path / "r.xlsx"
    write_excel_stream(iter(sr.files), str(out), summarize(sr), max_findings=3)
    wb = openpyxl.load_workbook(str(out))
    rows = list(wb["findings"].iter_rows(values_only=True))
    assert len(rows) == 1 + 3                       # 헤더 + 상한 3행
    summary = {str(r[0]): r[1] for r in wb["summary"].iter_rows(values_only=True) if r[0]}
    assert summary["노출(EXPOSED)"] == 5            # 집계는 생략과 무관하게 전체
    skipped = [v for k, v in summary.items() if "생략" in k]
    assert skipped == [2]


def test_excel_stream_consumes_generator_once(tmp_path):
    # 스트리밍 입력(1회성 제너레이터)으로도 4개 시트가 온전히 생성돼야 한다
    from pii_scanner.reporters.excel import write_excel_stream
    from pii_scanner.reporters.summary import summarize
    sr = _result()
    sr.files.append(FileResult(path="/x/locked.docx", encrypted=True))
    out = tmp_path / "r.xlsx"
    write_excel_stream(iter(sr.files), str(out), summarize(sr))
    wb = openpyxl.load_workbook(str(out))
    assert set(wb.sheetnames) == {"findings", "errors", "summary", "encrypted"}
    assert [r[0] for r in wb["encrypted"].iter_rows(values_only=True)][1:] == ["/x/locked.docx"]
    assert [r[0] for r in wb["errors"].iter_rows(values_only=True)][1:] == ["/x/bad.txt"]


def test_excel_summary_shows_corp_filtered(tmp_path):
    sr = _result()
    sr.files[0].corp_filtered = 7
    out = tmp_path / "r.xlsx"
    write_excel(sr, str(out))
    wb = openpyxl.load_workbook(str(out))
    rows = {str(r[0]): r[1] for r in wb["summary"].iter_rows(values_only=True) if r[0]}
    assert rows["법인등록번호 오탐 제거"] == 7


def test_excel_stream_keeps_key_pii_rows_when_cap_reached(tmp_path):
    """상한이 차도 고유식별정보 행은 남아야 한다 — 연락처류가 앞자리를 다 차지해도(T-017 후속).

    실측: 어느 대형 스캔에서 hit 의 94%가 이메일·유선전화였고, 파일 순서대로 채우자
    주민등록번호 노출 행의 79%가 엑셀에서 생략됐다.
    """
    from pii_scanner.reporters.excel import write_excel_stream
    from pii_scanner.reporters.summary import summarize
    sr = ScanResult()
    noisy = FileResult(path="/x/contacts.csv")
    for i in range(10):
        noisy.hits.append(PiiHit(PiiType.EMAIL, Status.EXPOSED, f"a{i}@ex***.com", 0, 12,
                                 Confidence.CONFIRMED, RiskLevel.MEDIUM))
    late = FileResult(path="/x/late.xlsx")
    late.hits.append(PiiHit(PiiType.RRN, Status.EXPOSED, "900101-1******", 0, 13,
                            Confidence.CONFIRMED, RiskLevel.CRITICAL))
    sr.files.extend([noisy, late])

    out = tmp_path / "r.xlsx"
    write_excel_stream(iter(sr.files), str(out), summarize(sr), max_findings=3)

    wb = openpyxl.load_workbook(str(out))
    rows = list(wb["findings"].iter_rows(values_only=True))[1:]
    assert "주민등록번호" in [r[2] for r in rows]      # 뒤에 나와도 반드시 수록
    assert len(rows) == 3                              # 상한 자체는 그대로 지킨다


def test_errors_sheet_separates_access_failure_from_extraction_failure(tmp_path):
    """감사에서 '폴더를 못 열었다'와 '파일 추출이 실패했다'는 의미가 다르다 — 열로 갈라야 한다."""
    import openpyxl
    from pii_scanner.core.models import FileResult, ScanResult
    from pii_scanner.reporters.excel import write_excel

    r = ScanResult()
    r.files.append(FileResult(path="/x.xls", error="추출 실패: XLRDError"))
    r.files.append(FileResult(path="/hr", error="접근 실패: 권한 없음", unreadable=True))
    out = tmp_path / "r.xlsx"
    write_excel(r, str(out))
    wb = openpyxl.load_workbook(str(out))
    rows = list(wb["errors"].iter_rows(values_only=True))
    assert rows[0] == ("파일경로", "사유", "구분")
    kinds = {row[0]: row[2] for row in rows[1:]}
    assert kinds == {"/x.xls": "파일 처리 실패", "/hr": "접근 실패"}
    summary = {row[0]: row[1] for row in wb["summary"].iter_rows(values_only=True) if row[0]}
    assert any("접근 실패" in k for k in summary), "summary 에 접근 실패 건수 행이 있어야 한다"


def test_summary_omits_access_failure_row_when_none(tmp_path):
    """0을 찍으면 순회 실패를 보고하지 않는 커넥터(Dropbox)에서 거짓 커버리지 보증이 된다."""
    import openpyxl
    from pii_scanner.core.models import FileResult, ScanResult
    from pii_scanner.reporters.excel import write_excel

    r = ScanResult()
    r.files.append(FileResult(path="/a.txt"))
    out = tmp_path / "r.xlsx"
    write_excel(r, str(out))
    wb = openpyxl.load_workbook(str(out))
    keys = [row[0] for row in wb["summary"].iter_rows(values_only=True) if row[0]]
    assert not any("접근 실패" in k for k in keys)


def test_illegal_control_chars_in_path_do_not_kill_the_report(tmp_path):
    """제어문자·surrogate 가 섞인 경로 때문에 리포트가 통째로 안 나오면 안 된다."""
    import openpyxl
    from pii_scanner.core.models import FileResult, ScanResult
    from pii_scanner.reporters.excel import write_excel

    r = ScanResult()
    r.files.append(FileResult(path="/bad\x01name\udcff", error="접근 실패: 권한 없음",
                              unreadable=True))
    out = tmp_path / "r.xlsx"
    write_excel(r, str(out))                      # 예외 없이 저장돼야 한다
    assert openpyxl.load_workbook(str(out))["errors"].max_row == 2


def test_summary_sheet_states_provenance_and_uncovered_types(tmp_path):
    """감사 산출물은 '언제·어느 버전으로·무엇을 검사했나'를 스스로 진술해야 한다."""
    import datetime, openpyxl
    from pii_scanner.config import ScanConfig
    from pii_scanner.core.models import FileResult, ScanResult
    from pii_scanner.reporters.excel import write_excel
    from pii_scanner.reporters.summary import stamp_scan_meta, summarize

    r = ScanResult(); r.files.append(FileResult(path="/a.txt"))
    s = stamp_scan_meta(summarize(r), ScanConfig(), now=datetime.datetime(2026, 8, 28, 9, 30))
    out = tmp_path / "r.xlsx"
    write_excel(r, str(out), summary=s)
    rows = {row[0]: row[1] for row in
            openpyxl.load_workbook(str(out))["summary"].iter_rows(values_only=True) if row[0]}
    assert rows["스캔 일시"] == "2026-08-28T09:30:00"
    assert rows["도구 버전"]
    assert "주민등록번호" in rows["검사한 종류(이번 실행)"]
    assert "외국인등록번호" in rows["⚠ 검사하지 않은 종류(이번 실행)"]
    assert "--enable foreign" in rows["⚠ 검사하지 않은 종류(이번 실행)"]
