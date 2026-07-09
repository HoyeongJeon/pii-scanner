import docx
import openpyxl
from pii_scanner.core.extractors.office import DocxExtractor, XlsxExtractor


def test_docx_extract(tmp_path):
    p = tmp_path / "a.docx"
    d = docx.Document()
    d.add_paragraph("홍길동 900101-1234568")
    d.save(str(p))
    assert "900101-1234568" in DocxExtractor().extract(str(p))


def test_xlsx_extract(tmp_path):
    p = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "홍길동"
    ws["B1"] = "010-1234-5678"
    wb.save(str(p))
    text = XlsxExtractor().extract(str(p))
    assert "010-1234-5678" in text


def test_xlsx_extract_located_labels_cells(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "이름"
    ws["C5"] = "kim@example.com"
    f = tmp_path / "m.xlsx"
    wb.save(str(f))

    text, locator = XlsxExtractor().extract_located(str(f))
    assert locator.label(text.index("이름")) == "Sheet1!A1"
    assert locator.label(text.index("kim@example.com")) == "Sheet1!C5"
    # extract()와 텍스트 동등 — 탐지 결과가 달라지지 않아야 한다
    assert XlsxExtractor().extract(str(f)) == text
