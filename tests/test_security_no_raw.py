import openpyxl
from pii_scanner.core.scanner import scan_paths
from pii_scanner.config import ScanConfig
from pii_scanner.core.extractors.base import Extractor
from pii_scanner.reporters.excel import write_excel
from pii_scanner.reporters.html import write_html

# 리포트에 절대 나오면 안 되는 "원문/민감부" (합성값).
# 주의: 휴대폰/유선 끝 4자리("5678"/"4567")는 승인된 마스킹 형식
# (`010-****-5678`, `02-***-4567`)의 정당한 잔여부이므로 유출이 아니다 — 목록에서 제외.
# 여기서는 "반드시 가려져야 하는 부분"과 "전체 원문 문자열"이 새는지를 검증한다.
RAW_SECRETS = [
    "900101-1234568",            # 주민번호 전체
    "1234568",                   # 주민번호 뒤 7자리(민감부) — 가려져야 함
    "010-1234-5678",             # 휴대폰 전체(가운데가 노출되면 안 됨)
    "1234-5678",                 # 휴대폰 가운데 노출 형태
    "M12345678",                 # 여권 전체
    "12345678",                  # 여권 본문 — 가려져야 함
    "11-12-345678-90",           # 운전면허 전체
    "345678",                    # 운전면허 일련번호 — 가려져야 함
    "02-123-4567",               # 유선전화 전체(가운데가 노출되면 안 됨)
    "hong.gildong@example.com",  # 이메일 전체
    "hong.gildong",              # 이메일 아이디 — 가려져야 함
]


_BODY = (
    "홍길동 900101-1234568\n"
    "메일 hong.gildong@example.com\n"
    "폰 010-1234-5678\n"
    "여권 M12345678\n"
    "면허 11-12-345678-90\n"
    "대표 02-123-4567\n"
)


def _make_src(tmp_path):
    """정상 파일 + '추출이 실패하는' 파일들.

    실패 경로를 안 만들면 errors 시트가 항상 비어서, 정작 유출이 일어나는 자리를
    이 회귀 테스트가 통과시킨다 — 실제로 그랬다. 서드파티 파싱 라이브러리는 실패한 파일의
    내용을 예외 메시지에 박는다(xlrd: "Expected BOF record; found b'900101-1'").
    아래 파일들은 전부 확장자와 내용이 어긋나 추출 단계에서 실패한다.
    """
    src = tmp_path / "src"; src.mkdir()
    (src / "a.txt").write_text(_BODY, encoding="utf-8")
    raw = _BODY.encode("utf-8")
    # 첫 바이트가 곧 유출량이다 — xlrd 는 파일 앞 8바이트를 예외 메시지에 그대로 박는다.
    # 주민번호로 시작하게 두어 '본문을 실으면 반드시 티가 나는' 픽스처로 만든다.
    (src / "payroll.xls").write_bytes(b"900101-1234568,hong.gildong,3000000\n")
    (src / "roster.xlsx").write_bytes(raw)         # zip 이 아님 → BadZipFile
    (src / "report.docx").write_bytes(raw)
    (src / "doc.hwpx").write_bytes(raw)
    (src / "scan.pdf").write_bytes(raw)
    return src


def test_excel_has_no_raw_pii(tmp_path):
    src = _make_src(tmp_path)
    result = scan_paths([str(src)], ScanConfig())
    out = tmp_path / "r.xlsx"
    write_excel(result, str(out))
    wb = openpyxl.load_workbook(str(out))
    blob = " ".join(
        str(c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    )
    assert wb["errors"].max_row > 1, "errors 시트가 비면 유출 지점을 검사하지 않은 것이다"
    for secret in RAW_SECRETS:
        assert secret not in blob, f"원문 유출: {secret}"


def test_html_has_no_raw_pii(tmp_path):
    src = _make_src(tmp_path)
    result = scan_paths([str(src)], ScanConfig())
    out = tmp_path / "r.html"
    write_html(result, str(out))
    blob = out.read_text(encoding="utf-8")
    for secret in RAW_SECRETS:
        assert secret not in blob, f"원문 유출: {secret}"


def test_state_jsonl_has_no_raw_pii(tmp_path):
    from pii_scanner.connectors.local_fs import LocalFsConnector
    from pii_scanner.connectors.state import ScanState
    from pii_scanner.core.scanner import scan

    src = _make_src(tmp_path)
    st = ScanState("sec", base_dir=str(tmp_path / "state"))
    scan(LocalFsConnector([str(src)]), ScanConfig(), state=st)
    blob = open(st.jsonl_path, encoding="utf-8").read()
    assert "추출 실패" in blob, "실패 경로가 만들어지지 않으면 이 테스트는 아무것도 검증하지 않는다"
    for secret in RAW_SECRETS:
        assert secret not in blob, f"상태파일 원문 유출: {secret}"


def test_parallel_state_jsonl_has_no_raw_pii(tmp_path, monkeypatch):
    # 병렬 경로(scan_parallel)도 상태파일에 마스킹 스니펫만 남겨야 한다.
    from pii_scanner.connectors.base import SourceFile, Connector
    from pii_scanner.connectors.state import ScanState
    from pii_scanner.core.scanner import scan_parallel
    import pii_scanner.core.scanner as scn

    src = _make_src(tmp_path)
    text = (src / "a.txt").read_text(encoding="utf-8")

    class _SF(SourceFile):
        def __init__(self, p): self.logical_path = p
        def __enter__(self): return self.logical_path
        def __exit__(self, *exc): pass
    class _Conn(Connector):
        def iter_files(self, config):
            yield _SF("/d/a.txt")
    class _Ext(Extractor):
        def extract(self, local): return text
    monkeypatch.setattr(scn, "get_extractor", lambda path: _Ext())

    st = ScanState("secpar", base_dir=str(tmp_path / "state"))
    scan_parallel(_Conn(), ScanConfig(), state=st, max_workers=4)
    blob = open(st.jsonl_path, encoding="utf-8").read()
    for secret in RAW_SECRETS:
        assert secret not in blob, f"병렬 상태파일 원문 유출: {secret}"


def test_dropbox_temp_is_outside_any_sync_folder():
    import os, tempfile
    import pytest
    pytest.importorskip("dropbox")
    from unittest.mock import MagicMock
    from pii_scanner.connectors.dropbox_conn import DropboxSourceFile

    dbx = MagicMock()
    dbx.files_download_to_file.side_effect = lambda local, path: open(local, "wb").close()
    sf = DropboxSourceFile(dbx, "/팀/a.txt", "/팀/a.txt")
    with sf as local:
        # 시스템 temp(동기화 폴더 밖). 'Dropbox' 경로 조각이 들어가면 안 됨.
        assert local.startswith(tempfile.gettempdir())
        assert "Dropbox" not in local and "dropbox" not in os.path.dirname(local).lower()


def _error_strings(out_path):
    import openpyxl
    ws = openpyxl.load_workbook(str(out_path))["errors"]
    return [str(row[1]) for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]


def test_error_reasons_never_carry_third_party_message_bodies(tmp_path):
    """사유는 예외 '타입'만. 파싱 라이브러리는 실패한 파일의 내용을 메시지에 박는다 —
    xlrd 의 getbof 는 파일 앞 8바이트를 넣고, 그 자리에 주민번호·이메일이 올 수 있다.
    유출량은 라이브러리마다 다르므로 고정 문자열 목록이 아니라 '모양'으로 막는다."""
    import re

    src = _make_src(tmp_path)
    result = scan_paths([str(src)], ScanConfig())
    out = tmp_path / "r.xlsx"
    write_excel(result, str(out))

    reasons = _error_strings(out)
    assert reasons, "실패 파일이 없으면 이 테스트는 아무것도 검증하지 않는다"
    for reason in reasons:
        assert not re.search(r"\d{4}", reason), f"사유에 숫자열이 실렸다(원문 유출 의심): {reason}"
        assert "@" not in reason, f"사유에 이메일 흔적이 실렸다: {reason}"
        assert reason.split(":")[0] in ("추출 실패", "탐지 일부 실패", "접근 실패"), reason
