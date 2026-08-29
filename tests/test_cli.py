import os
import pytest
from pii_scanner.cli import main


def test_cli_scans_and_writes_reports(tmp_path):
    src = tmp_path / "src"; src.mkdir()
    (src / "a.txt").write_text("홍길동 900101-1234568", encoding="utf-8")
    outdir = tmp_path / "out"
    rc = main([str(src), "--out", str(outdir)])
    assert rc == 0
    assert os.path.exists(outdir / "report.xlsx")
    assert os.path.exists(outdir / "report.html")


def test_cli_disable_detector(tmp_path):
    src = tmp_path / "src"; src.mkdir()
    (src / "a.txt").write_text("ho@example.com", encoding="utf-8")
    outdir = tmp_path / "out"
    rc = main([str(src), "--out", str(outdir), "--disable", "email"])
    assert rc == 0


def test_cli_unknown_disable_key_errors(tmp_path):
    src = tmp_path / "src"; src.mkdir()
    (src / "a.txt").write_text("x", encoding="utf-8")
    # 오타("emial")는 조용히 무시되지 않고 에러로 종료해야 함
    with pytest.raises(SystemExit):
        main([str(src), "--out", str(tmp_path / "out"), "--disable", "emial"])


def test_cli_foreign_off_by_default(tmp_path):
    import openpyxl
    src = tmp_path / "src"; src.mkdir()
    (src / "a.txt").write_text("외국인등록 900101-5234561 끝", encoding="utf-8")
    out = tmp_path / "out"
    rc = main([str(src), "--out", str(out)])          # 기본: 외국인 비활성(후순위)
    assert rc == 0
    wb = openpyxl.load_workbook(str(out / "report.xlsx"))
    blob = " ".join(str(c.value) for r in wb["findings"].iter_rows()
                    for c in r if c.value)
    assert "외국인등록번호" not in blob


def test_cli_enable_foreign_detects(tmp_path):
    import openpyxl
    src = tmp_path / "src"; src.mkdir()
    (src / "a.txt").write_text("외국인등록 900101-5234561 끝", encoding="utf-8")
    out = tmp_path / "out"
    rc = main([str(src), "--out", str(out), "--enable", "foreign"])   # 명시적으로 켜기
    assert rc == 0
    wb = openpyxl.load_workbook(str(out / "report.xlsx"))
    blob = " ".join(str(c.value) for r in wb["findings"].iter_rows()
                    for c in r if c.value)
    assert "외국인등록번호" in blob


def test_cli_unknown_enable_key_errors(tmp_path):
    src = tmp_path / "src"; src.mkdir()
    (src / "a.txt").write_text("x", encoding="utf-8")
    with pytest.raises(SystemExit):
        main([str(src), "--out", str(tmp_path / "o"), "--enable", "nope"])


def test_cli_exits_nonzero_and_reports_when_root_is_unreadable(tmp_path, capsys):
    """경로 오타·권한 차단이 '깨끗함 + exit 0' 으로 보고되면 감사가 무효가 된다."""
    import openpyxl
    from pii_scanner import cli

    out = tmp_path / "out"
    rc = cli.main([str(tmp_path / "없는부서폴더"), "--out", str(out)])
    assert rc == cli.EXIT_ACCESS_ERROR
    assert rc != 2, "2는 argparse 사용법 오류용 — 자동화가 구분할 수 있어야 한다"
    err = capsys.readouterr().err
    assert "읽지 못해" in err

    wb = openpyxl.load_workbook(str(out / "report.xlsx"))
    rows = [r for r in wb["errors"].iter_rows(values_only=True)][1:]
    assert [(r[0], r[2]) for r in rows] == [(str(tmp_path / "없는부서폴더"), "접근 실패")]


def test_cli_exit_zero_when_everything_readable(tmp_path):
    from pii_scanner import cli

    (tmp_path / "a.txt").write_text("내용 없음", encoding="utf-8")
    assert cli.main([str(tmp_path), "--out", str(tmp_path / "o")]) == 0
