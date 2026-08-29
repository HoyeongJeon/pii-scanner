from pii_scanner.core.scanner import process_file
from pii_scanner.core.detectors import build_detectors
from pii_scanner.config import ScanConfig
from pii_scanner.connectors.base import SourceFile
from pii_scanner.core.extractors import EncryptedFileError, UnsupportedFormat
from pii_scanner.core.extractors.base import Extractor


class _FakeSF(SourceFile):
    def __init__(self, path):
        self.logical_path = path
        self.exited = False
    def __enter__(self):
        return self.logical_path
    def __exit__(self, *exc):
        self.exited = True


def _detectors():
    return build_detectors(ScanConfig())


def _patch_ext(monkeypatch, behavior):
    import pii_scanner.core.scanner as scn
    class _Ext(Extractor):
        def __init__(self, path): self.path = path
        def extract(self, local):
            v = behavior[self.path]
            if isinstance(v, Exception):
                raise v
            return v
    monkeypatch.setattr(scn, "get_extractor", lambda path: _Ext(path))


def test_process_file_detects_and_cleans(monkeypatch):
    _patch_ext(monkeypatch, {"/d/a.txt": "주민 900101-1234568"})
    sf = _FakeSF("/d/a.txt")
    fr = process_file(sf, _detectors())
    assert fr.path == "/d/a.txt"
    assert fr.hits and fr.error is None and fr.encrypted is False
    assert sf.exited is True


def test_process_file_records_extraction_error(monkeypatch):
    _patch_ext(monkeypatch, {"/d/bad.docx": ValueError("boom")})
    fr = process_file(_FakeSF("/d/bad.docx"), _detectors())
    assert fr.error is not None and fr.encrypted is False and fr.hits == []


def test_process_file_marks_encrypted(monkeypatch):
    _patch_ext(monkeypatch, {"/d/lock.docx": EncryptedFileError("locked")})
    fr = process_file(_FakeSF("/d/lock.docx"), _detectors())
    assert fr.encrypted is True and fr.error is None


def test_process_file_returns_none_for_unsupported(monkeypatch):
    _patch_ext(monkeypatch, {"/d/x.bin": UnsupportedFormat("nope")})
    assert process_file(_FakeSF("/d/x.bin"), _detectors()) is None


class _Boom:
    def find(self, text):
        raise RuntimeError("detector boom")


def test_process_file_isolates_detector_exception(monkeypatch):
    # 탐지기가 예외를 던져도 process_file 은 밖으로 던지지 않고 fr.error 로 격리한다
    # (워커 스레드에서 호출되므로 — scan_parallel 의 fut.result() 를 죽이면 안 됨).
    _patch_ext(monkeypatch, {"/d/a.txt": "주민 900101-1234568"})
    fr = process_file(_FakeSF("/d/a.txt"), [_Boom()])
    assert fr is not None and fr.error is not None and "탐지" in fr.error
    assert fr.partial_detection == ["_Boom(RuntimeError)"]


def test_one_throwing_detector_does_not_erase_the_other_detectors_hits(monkeypatch):
    """try 가 탐지기 루프 '밖'에 있으면 하나가 던졌을 때 그 파일이 통째로 0건이 된다 —
    실제로 3건 나오던 파일이 hits=[] 로 바뀌었다."""
    from pii_scanner.core.detectors import build_detectors
    from pii_scanner.config import ScanConfig

    text = "주민 900101-1234568 / 폰 010-1234-5678 / 메일 a@example.com"
    _patch_ext(monkeypatch, {"/d/a.txt": text})
    good = build_detectors(ScanConfig())
    assert len(process_file(_FakeSF("/d/a.txt"), good).hits) == 3      # 대조군

    mixed = good[:1] + [_Boom()] + good[1:]
    fr = process_file(_FakeSF("/d/a.txt"), mixed)
    assert len(fr.hits) == 3, "던지는 탐지기 하나가 나머지 결과를 지우면 안 된다"
    assert fr.partial_detection == ["_Boom(RuntimeError)"]
    assert fr.error and "일부 실패" in fr.error   # 그래도 '깨끗함'으로 보이면 안 된다


def test_process_file_sets_cell_location_for_xlsx(tmp_path):
    # 실제 XlsxExtractor 경유(monkeypatch 없음) — 셀 좌표가 hit에 실리는지
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["C5"] = "kim@example.com"
    f = tmp_path / "members.xlsx"
    wb.save(str(f))

    fr = process_file(_FakeSF(str(f)), _detectors())
    assert fr.hits and fr.error is None
    assert fr.hits[0].location == "Sheet1!C5"


def test_process_file_sets_line_location_for_text(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("첫줄\n주민 900101-1234568\n", encoding="utf-8")
    fr = process_file(_FakeSF(str(f)), _detectors())
    assert fr.hits and fr.hits[0].location == "L2"


def test_process_file_filters_corp_regno_column(tmp_path):
    # 법인번호(양쪽 체크섬 통과) 4건이 B열에 몰림 + D1에 이메일 → RRN 드롭, 이메일 생존
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i, num in enumerate(["110111-1000481", "110111-1000572",
                             "110111-1000621", "110111-1000663"], start=1):
        ws[f"B{i}"] = num
    ws["D1"] = "kim@example.com"
    f = tmp_path / "corp.xlsx"
    wb.save(str(f))

    fr = process_file(_FakeSF(str(f)), _detectors())
    assert fr.error is None
    types = {h.pii_type.value for h in fr.hits}
    assert "주민등록번호" not in types          # 법인번호 열 오탐 제거됨
    assert "이메일" in types                     # 진짜 PII 생존
    assert fr.corp_filtered == 4
